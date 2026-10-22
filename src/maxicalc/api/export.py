"""B2B estimate (смета) export: XLSX (openpyxl) and PDF (fpdf2, needs a Cyrillic TTF font)."""
import io
import os
from decimal import Decimal
from pathlib import Path

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from maxicalc.api.schemas import EstimateOut

FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
)
KIND_RU = {"property": "Смета на объект", "room": "Смета на помещение",
           "product": "Расчёт количества по карточке товара"}


class PdfUnavailable(RuntimeError):
    pass


def find_font(configured: str | None = None) -> str | None:
    for p in (configured, os.environ.get("PDF_FONT_PATH"), *FONT_CANDIDATES):
        if p and Path(p).is_file():
            return p
    return None


def _num(x: Decimal | None) -> float | None:
    return None if x is None else float(x)


def build_xlsx(est: EstimateOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Смета"
    bold = Font(bold=True)
    ws.append([KIND_RU.get(est.kind, "Смета")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Номер сметы: {est.estimate_id}", f"Версия каталога: {est.catalog_version}"])
    ws.append([])
    ws.append(["№", "Артикул", "Наименование", "Кол-во", "Ед.", "Цена, ₽", "Сумма, ₽",
               "Обоснование"])
    head_row = ws.max_row
    for c in ws[head_row]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
    first = ws.max_row + 1
    for i, ln in enumerate(est.basket.lines, start=1):
        ws.append([i, ln.sku, ln.name, ln.qty_packs, ln.unit, _num(ln.unit_price),
                   _num(ln.line_total), ln.reason])
    last = ws.max_row
    ws.append([])
    ws.append(["", "", "Итого материалы", "", "", "", _num(est.basket.materials_total)])
    ws.cell(row=ws.max_row, column=3).font = bold
    ws.cell(row=ws.max_row, column=7).font = bold
    ws.append(["", "", "Стоимость работ не рассчитана: ставки работ не заданы"])
    for r in range(first, last + 2):
        for col in (6, 7):
            ws.cell(row=r, column=col).number_format = "#,##0.00"
    for col, w in zip(range(1, 9), (5, 14, 52, 9, 7, 12, 14, 48), strict=True):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    calc = wb.create_sheet("Расчёт")
    calc.append(["Артикул", "Наименование", "Ход расчёта"])
    for c in calc[1]:
        c.font = bold
    for ln in est.basket.lines:
        for k, line in enumerate(ln.explain):
            calc.append([ln.sku if k == 0 else "", ln.name if k == 0 else "", line])
    calc.column_dimensions["A"].width = 14
    calc.column_dimensions["B"].width = 40
    calc.column_dimensions["C"].width = 110

    notes = wb.create_sheet("Допущения")
    notes.append(["Допущения"])
    notes["A1"].font = bold
    for a in est.assumptions or ["—"]:
        notes.append([a])
    notes.append([])
    notes.append(["Предупреждения"])
    notes.cell(row=notes.max_row, column=1).font = bold
    for w in est.warnings or ["—"]:
        notes.append([w])
    for e in est.electrical:
        notes.append([f"Электрика, {e.room}: розетки {e.sockets}, выключатели {e.switches}, "
                      f"светильники {e.light_points}, кабель {e.cable_m} м, автоматы "
                      f"{e.breakers} ({e.note})"])
    notes.column_dimensions["A"].width = 140

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(est: EstimateOut, font_path: str | None = None) -> bytes:
    font = find_font(font_path)
    if font is None:
        raise PdfUnavailable(
            "PDF недоступен: не найден шрифт с кириллицей. Укажите PDF_FONT_PATH "
            "(например, DejaVuSans.ttf)")
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("Body", "", font)
    pdf.add_page()
    pdf.set_font("Body", size=14)
    pdf.cell(0, 8, KIND_RU.get(est.kind, "Смета"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Body", size=8)
    pdf.cell(0, 5, f"Номер сметы: {est.estimate_id}   Версия каталога: {est.catalog_version}",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Body", size=8)
    head = FontFace(emphasis=None, fill_color=(230, 230, 230))  # only the regular TTF is loaded
    with pdf.table(col_widths=(7, 22, 71, 14, 36, 40), headings_style=head,
                   text_align=("CENTER", "LEFT", "LEFT", "RIGHT", "RIGHT", "RIGHT")) as table:
        row = table.row()
        for h in ("№", "Артикул", "Наименование", "Кол-во", "Цена, ₽", "Сумма, ₽"):
            row.cell(h)
        for i, ln in enumerate(est.basket.lines, start=1):
            row = table.row()
            row.cell(str(i))
            row.cell(ln.sku)
            row.cell(ln.name)
            row.cell(f"{ln.qty_packs} {ln.unit}")
            row.cell("—" if ln.unit_price is None else f"{ln.unit_price:,.2f}")
            row.cell(f"{ln.line_total:,.2f}")
    pdf.ln(3)
    pdf.set_font("Body", size=10)
    pdf.cell(0, 6, f"Итого материалы: {est.basket.materials_total:,.2f} ₽",
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Body", size=8)
    pdf.multi_cell(0, 4, "Стоимость работ не рассчитана: ставки работ не заданы.",
                   new_x="LMARGIN", new_y="NEXT")
    if est.assumptions:
        pdf.ln(2)
        pdf.cell(0, 5, "Допущения:", new_x="LMARGIN", new_y="NEXT")
        for a in est.assumptions:
            pdf.multi_cell(0, 4, f"• {a}", new_x="LMARGIN", new_y="NEXT")
    if est.warnings:
        pdf.ln(2)
        pdf.cell(0, 5, "Предупреждения:", new_x="LMARGIN", new_y="NEXT")
        for w in est.warnings:
            pdf.multi_cell(0, 4, f"• {w}", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
