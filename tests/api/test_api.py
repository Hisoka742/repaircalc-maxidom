import io
import time
from decimal import Decimal as D

import openpyxl
import pytest
from fastapi.testclient import TestClient

from maxicalc.api.export import find_font
from maxicalc.api.main import create_app
from maxicalc.api.state import AppState
from maxicalc.catalog.schema import Product
from maxicalc.config import Settings
from maxicalc.ml.match import Matcher
from maxicalc.recommend.build import to_entries
from maxicalc.recommend.select import CatalogEntry


def e(sku: str, calc: str, price: str, attrs: dict[str, object],
      props: set[str] | None = None, name: str | None = None) -> CatalogEntry:
    return CatalogEntry(
        Product(sku=sku, name=name or f"Товар {sku}", price=D(price)), work=set(),
        rooms={"any"}, props=props or set(), calc=calc, status="verified",
        attrs={k: (v if isinstance(v, int) else D(str(v))) for k, v in attrs.items()})


ENTRIES = [
    e("P1", "paint", "1000", {"spread_m2_per_l": 10, "pack_volume_l": "2.5"}, {"washable"}),
    e("R1", "primer", "500", {"consumption_l_per_m2": "0.15", "pack_volume_l": 5}),
    e("L1", "laminate", "1500", {"m2_per_pack": "2.131"}, {"wear_class_33"}),
    e("U1", "underlay", "300", {"m2_per_pack": 10}),
    e("S1", "skirting", "200", {"piece_len_m": "2.5"}),
    e("W1", "wallpaper", "1500", {"roll_width_m": "1.06", "roll_length_m": "10.05"}),
    e("D1", "drywall", "400", {}, {"moisture_resistant"}),
    e("D2", "drywall", "300", {}),
    e("E1", "electrical", "100", {}),
]


def make_client(entries: list[CatalogEntry] | None = None, **cfg: object) -> TestClient:
    settings = Settings(rate_limit_per_min=0, log_json=True, **cfg)  # type: ignore[arg-type]
    state = AppState(entries if entries is not None else list(ENTRIES), settings)
    return TestClient(create_app(state, settings))


@pytest.fixture()
def client() -> TestClient:
    return make_client()


ROOM_BODY = {"room": {"length_m": "4", "width_m": "3", "height_m": "2.5"},
             "room_type": "living", "renovation_class": "cosmetic"}


# ---------- happy paths ----------
def test_health_reports_catalog(client: TestClient) -> None:
    j = client.get("/health").json()
    assert j["status"] == "ok" and j["catalog_size"] == len(ENTRIES) and j["cache"] == "memory"
    assert len(j["catalog_version"]) == 12


def test_room_calc_matches_hand_computed_basket(client: TestClient) -> None:
    r = client.post("/v1/calc/room", json=ROOM_BODY)
    assert r.status_code == 200
    j = r.json()
    q = {ln["sku"]: ln for ln in j["basket"]["lines"]}
    assert (q["P1"]["qty_packs"], q["R1"]["qty_packs"], q["L1"]["qty_packs"],
            q["U1"]["qty_packs"], q["S1"]["qty_packs"]) == (4, 2, 6, 2, 6)
    assert D(j["basket"]["materials_total"]) == D("15800")
    assert j["labor"] is not None
    assert D(j["labor"]["low"]) < D(j["labor"]["expected"]) < D(j["labor"]["high"])
    assert all(ln["explain"] and ln["reason"] for ln in j["basket"]["lines"])
    assert j["kind"] == "room" and len(j["estimate_id"]) == 16


def test_property_calc(client: TestClient) -> None:
    r = client.post("/v1/calc/property", json={
        "property_type": "apartment", "total_area_m2": "60", "renovation_class": "cosmetic"})
    assert r.status_code == 200
    j = r.json()
    assert sum(D(x["area_m2"]) for x in j["rooms"]) == D("60")
    assert j["basket"]["lines"]


def test_estimates_are_idempotent_and_cached(client: TestClient) -> None:
    a = client.post("/v1/calc/room", json=ROOM_BODY).json()
    reordered = dict(reversed(list(ROOM_BODY.items())))
    b = client.post("/v1/calc/room", json=reordered).json()
    assert a["estimate_id"] == b["estimate_id"]
    assert a["cached"] is False and b["cached"] is True
    assert a["basket"] == b["basket"]
    other = client.post("/v1/calc/room", json={**ROOM_BODY, "renovation_class": "major"}).json()
    assert other["estimate_id"] != a["estimate_id"]


def test_estimate_id_changes_with_catalog_version() -> None:
    a = make_client().post("/v1/calc/room", json=ROOM_BODY).json()
    changed = [*ENTRIES[:-1], e("E9", "electrical", "1", {})]
    b = make_client(changed).post("/v1/calc/room", json=ROOM_BODY).json()
    assert a["estimate_id"] != b["estimate_id"] and a["catalog_version"] != b["catalog_version"]


def test_product_card_quantity(client: TestClient) -> None:
    body = {"room": {"length_m": "4", "width_m": "3", "height_m": "2.5",
                     "openings": [{"kind": "door", "width_m": "0.9", "height_m": "2.05",
                                   "wall": 0}]}, "room_type": "living"}
    r = client.post("/v1/products/W1/quantity", json=body)
    assert r.status_code == 200
    q = {ln["sku"]: ln for ln in r.json()["basket"]["lines"]}
    assert q["W1"]["qty_packs"] == 5 and q["W1"]["unit"] == "рулон"
    assert r.json()["kind"] == "product"


def test_product_card_errors(client: TestClient) -> None:
    body = {"room": {"length_m": "3", "width_m": "3"}, "room_type": "kitchen"}
    r = client.post("/v1/products/NOPE/quantity", json=body)
    assert r.status_code == 404 and r.json()["error"]["code"] == "http_404"
    r = client.post("/v1/products/E1/quantity", json=body)     # electrical: no quantity calc
    assert r.status_code == 422 and "калькулятор" in r.json()["error"]["message"]


def test_l_shaped_room_by_vertices(client: TestClient) -> None:
    body = {"room": {"vertices": [["0", "0"], ["4", "0"], ["4", "2"], ["2", "2"], ["2", "4"],
                                  ["0", "4"]], "height_m": "2.5"},
            "room_type": "living", "renovation_class": "cosmetic"}
    r = client.post("/v1/calc/room", json=body)
    assert r.status_code == 200 and D(r.json()["rooms"][0]["area_m2"]) == 12


def test_recommend_filters_by_room(client: TestClient) -> None:
    j = client.post("/v1/recommend", json={"room_type": "bathroom", "calc": "drywall"}).json()
    assert [i["sku"] for i in j["items"]] == ["D1"]
    assert any("влагостойк" in x["reason"] for x in j["excluded"]) and j["excluded_total"] == 1
    assert "влагостойк" in " ".join(j["items"][0]["reasons"])
    empty = client.post("/v1/recommend", json={"room_type": "bathroom", "calc": "tile"}).json()
    assert empty["items"] == [] and empty["warnings"]


# ---------- Russian validation errors ----------
def test_validation_errors_are_russian_and_field_specific(client: TestClient) -> None:
    r = client.post("/v1/calc/property", json={
        "property_type": "castle", "total_area_m2": "-5", "renovation_class": "cosmetic",
        "bogus": 1})
    assert r.status_code == 422
    err = r.json()["error"]
    assert err["code"] == "validation_error"
    by_field = {d["field"]: d["message"] for d in err["details"]}
    assert "больше" in by_field["total_area_m2"]
    assert "Допустимые значения" in by_field["property_type"]
    assert by_field["bogus"] == "Лишнее поле"
    missing = client.post("/v1/calc/property", json={}).json()["error"]["details"]
    assert {d["field"] for d in missing} == {"property_type", "total_area_m2",
                                              "renovation_class"}
    assert all(d["message"] == "Обязательное поле" for d in missing)


def test_geometry_needs_one_shape(client: TestClient) -> None:
    body = {"room": {"height_m": "2.5"}, "room_type": "living", "renovation_class": "cosmetic"}
    err = client.post("/v1/calc/room", json=body).json()["error"]
    assert "Укажите длину и ширину" in err["details"][0]["message"]


def test_bad_json_and_unknown_route(client: TestClient) -> None:
    r = client.post("/v1/calc/room", content=b"{oops", headers={"content-type": "application/json"})
    assert r.status_code == 422 and "JSON" in r.json()["error"]["details"][0]["message"]
    r = client.get("/v1/nothing")
    assert r.status_code == 404 and r.json()["error"]["message"] == "Не найдено"


def test_engine_errors_become_russian_422(client: TestClient) -> None:
    body = {**ROOM_BODY, "room": {**ROOM_BODY["room"], "openings": [  # type: ignore[dict-item]
        {"kind": "door", "width_m": "4.5", "height_m": "2", "wall": 0}]}}
    r = client.post("/v1/calc/room", json=body)
    assert r.status_code == 422
    assert r.json()["error"] == {"code": "invalid_request", "message": "Проёмы шире стены",
                                 "details": []}


# ---------- observability / limits ----------
def test_request_id_and_metrics(client: TestClient) -> None:
    r = client.get("/health", headers={"X-Request-ID": "abc123"})
    assert r.headers["x-request-id"] == "abc123"
    assert client.get("/health").headers["x-request-id"]
    client.post("/v1/calc/room", json=ROOM_BODY)
    text = client.get("/metrics").text
    assert 'http_requests_total{method="POST",route="/v1/calc/room",status="200"}' in text
    assert "http_request_duration_seconds_bucket" in text and "catalog_products" in text


def test_rate_limit_returns_429_in_russian_and_exempts_health() -> None:
    settings = Settings(rate_limit_per_min=60, rate_limit_burst=3)
    state = AppState(list(ENTRIES), settings)
    c = TestClient(create_app(state, settings))
    codes = [c.post("/v1/recommend", json={"room_type": "living", "calc": "paint"}).status_code
             for _ in range(5)]
    assert codes[:3] == [200, 200, 200] and codes[3] == 429
    r = c.post("/v1/recommend", json={"room_type": "living", "calc": "paint"})
    assert r.status_code == 429 and int(r.headers["retry-after"]) >= 1
    assert "Слишком много запросов" in r.json()["error"]["message"]
    assert c.get("/health").status_code == 200


# ---------- matching, jobs, review ----------
PRODUCTS = [
    {"sku": "a", "name": "Обои флизелиновые 1,06х10,05 м", "category_path": ["Обои"]},
    {"sku": "b", "name": "Клей плиточный Ceresit CM 11, 25 кг",
     "category_path": ["Плитка", "Клей для плитки"]},
    {"sku": "c", "name": "Ceresit CM 11 PRO"},
]


def test_match_sync(client: TestClient) -> None:
    j = client.post("/v1/catalog/match", json={"products": PRODUCTS}).json()
    assert j["status"] == "done"
    by = {x["sku"]: x for x in j["results"]}
    assert by["a"]["status"] == "auto" and by["a"]["calc_template"]["label"] == "wallpaper"
    assert by["a"]["attributes"]["roll_width_m"] == "1.06"
    assert by["b"]["calc_template"]["label"] == "tile_adhesive"
    assert by["c"]["status"] == "review"


def test_match_large_batch_runs_as_background_job() -> None:
    c = make_client(sync_match_limit=2)
    r = c.post("/v1/catalog/match", json={"products": PRODUCTS * 2})
    assert r.status_code == 202
    job_id = r.json()["job_id"]
    for _ in range(100):
        j = c.get(f"/v1/catalog/jobs/{job_id}").json()
        if j["status"] != "running":
            break
        time.sleep(0.05)
    assert j["status"] == "done" and j["done"] == j["total"] == 6 and len(j["results"]) == 6
    assert c.get("/v1/catalog/jobs/nope").status_code == 404


def test_review_loop_updates_live_catalog() -> None:
    prods = [Product(sku="x1", name="Ceresit CM 11 PRO", price=D("600")),
             Product(sku="x2", name="Обои флизелиновые 1,06х10,05 м", category_path=["Обои"],
                     price=D("900"))]
    results = Matcher().match(prods)
    settings = Settings(rate_limit_per_min=0)
    state = AppState(to_entries(prods, results), settings, results=results)
    c = TestClient(create_app(state, settings))
    q = c.get("/v1/catalog/review-queue").json()
    assert [i["sku"] for i in q] == ["x1"]
    v0 = c.get("/health").json()["catalog_version"]
    ok = c.post("/v1/catalog/review/x1", json={
        "work_categories": ["wall_tile"], "room_types": ["any"], "properties": [],
        "calc_template": "tile_adhesive", "reviewer": "ivan"})
    assert ok.status_code == 200 and ok.json()["catalog_version"] != v0
    assert c.get("/v1/catalog/review-queue").json() == []
    rec = c.post("/v1/recommend", json={"room_type": "living", "calc": "tile_adhesive"}).json()
    assert [i["sku"] for i in rec["items"]] == ["x1"]      # now served from the live catalog
    assert c.post("/v1/catalog/review/zzz", json={
        "work_categories": [], "calc_template": "none", "reviewer": "ivan"}).status_code == 404
    bad = c.post("/v1/catalog/review/x2", json={
        "work_categories": ["nonsense"], "calc_template": "none", "reviewer": "ivan"})
    assert bad.status_code == 422 and "Неизвестный вид работ" in bad.json()["error"]["message"]


# ---------- export ----------
def test_export_xlsx(client: TestClient) -> None:
    eid = client.post("/v1/calc/room", json=ROOM_BODY).json()["estimate_id"]
    r = client.post(f"/v1/estimate/{eid}/export?format=xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert f"smeta-{eid}.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Смета", "Расчёт", "Допущения"]
    ws = wb["Смета"]
    cells = [c for row in ws.iter_rows(values_only=True) for c in row]
    assert "P1" in cells and "Итого материалы" in cells and 15800.0 in cells
    assert client.post("/v1/estimate/deadbeef/export").status_code == 404


def test_export_pdf(client: TestClient) -> None:
    eid = client.post("/v1/calc/room", json=ROOM_BODY).json()["estimate_id"]
    r = client.post(f"/v1/estimate/{eid}/export?format=pdf")
    if find_font() is None:
        assert r.status_code == 501 and "шрифт" in r.json()["error"]["message"]
    else:
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert client.post(f"/v1/estimate/{eid}/export?format=docx").status_code == 422


def test_openapi_documents_all_endpoints(client: TestClient) -> None:
    # the case's paths are /api/v1/...; the original /v1/... are hidden aliases (A2)
    paths = client.get("/openapi.json").json()["paths"]
    for p in ["/api/v1/calc/property", "/api/v1/calc/room", "/api/v1/products/{sku}/quantity",
              "/api/v1/recommend", "/api/v1/catalog/match", "/api/v1/catalog/review-queue",
              "/api/v1/catalog/review/{sku}", "/api/v1/estimate/{estimate_id}/export",
              "/health"]:
        assert p in paths, p
    ex = paths["/api/v1/calc/room"]["post"]["requestBody"]["content"]["application/json"]["schema"]
    assert ex  # request schema published
